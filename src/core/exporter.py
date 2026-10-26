"""
Export module for generating Excel (.xlsx) and CSV reports.
Uses openpyxl for clean formatting, styling, and multiple sheets.
"""
from datetime import datetime
import os
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import AggregationSummary, GroupSummaryRow


def export_to_excel(
    summary: AggregationSummary,
    file_name: str,
    output_path: str,
    grouping_title: str = "Категория / Группа",
) -> str:
    """
    Export aggregation summary and list of missing elements to an Excel workbook.
    """
    wb = openpyxl.Workbook()
    
    # Setup styles
    font_title = Font(name="Segoe UI", size=14, bold=True, color="1E293B")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    font_meta_label = Font(name="Segoe UI", size=10, bold=True, color="334155")
    font_meta_val = Font(name="Segoe UI", size=10, color="0F172A")
    
    font_th = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fill_th = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    
    font_tr = Font(name="Segoe UI", size=10, color="0F172A")
    fill_alt = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    font_total = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    
    border_thin = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
    border_total_top = Border(top=Side(border_style="thin", color="0F172A"), bottom=Side(border_style="double", color="0F172A"))
    
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    # -------------------------------------------------------------
    # SHEET 1: Сводная таблица
    # -------------------------------------------------------------
    ws_summary = wb.active
    ws_summary.title = "Сводная таблица"
    ws_summary.views.sheetView[0].showGridLines = True

    # Title block
    ws_summary["A1"] = "Отчет по анализу атрибутов BIM-модели IFC"
    ws_summary["A1"].font = font_title
    ws_summary["A2"] = f"Сформировано: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws_summary["A2"].font = font_subtitle

    # Metadata block
    meta_data = [
        ("Файл модели:", file_name),
        ("Атрибут / Свойство:", summary.attribute.full_path),
        ("Тип свойства:", f"{summary.attribute.source.value} ({summary.attribute.data_type.value})"),
        ("Всего элементов:", f"{summary.total_elements:,}"),
        ("С атрибутом:", f"{summary.present_count:,} ({summary.present_percent}%)"),
        ("БЕЗ атрибута:", f"{summary.missing_count:,} ({summary.missing_percent}%)"),
    ]
    if summary.total_sum is not None:
        meta_data.append(("ИТОГОВАЯ СУММА:", f"{summary.total_sum:,.2f} {summary.attribute.unit}"))
    if summary.average is not None:
        meta_data.append(("Среднее значение:", f"{summary.average:,.2f} {summary.attribute.unit}"))

    row_cursor = 4
    for label, val in meta_data:
        ws_summary.cell(row=row_cursor, column=1, value=label).font = font_meta_label
        ws_summary.cell(row=row_cursor, column=2, value=val).font = font_meta_val
        row_cursor += 1

    row_cursor += 1  # blank line

    # Table Headers
    headers = [
        grouping_title,
        "Элементов всего",
        "С атрибутом (шт.)",
        "С атрибутом (%)",
        "Без атрибута (шт.)",
        "Без атрибута (%)",
    ]
    is_numeric = summary.total_sum is not None
    if is_numeric:
        headers.extend(["Сумма", "Среднее", "Минимум", "Максимум"])

    for col_idx, h_text in enumerate(headers, start=1):
        c = ws_summary.cell(row=row_cursor, column=col_idx, value=h_text)
        c.font = font_th
        c.fill = fill_th
        c.alignment = align_center
        c.border = cell_border

    start_table_row = row_cursor + 1
    row_cursor += 1

    # Populate data rows
    for r_idx, row in enumerate(summary.rows):
        is_even = (r_idx % 2 == 0)
        c_fill = None if is_even else fill_alt

        cols = [
            (row.group_name, align_left, "@"),
            (row.total_elements, align_right, "#,##0"),
            (row.present_count, align_right, "#,##0"),
            (row.present_percent / 100.0, align_right, "0.0%"),
            (row.missing_count, align_right, "#,##0"),
            (row.missing_percent / 100.0, align_right, "0.0%"),
        ]
        if is_numeric:
            cols.extend([
                (row.total_sum if row.total_sum is not None else "-", align_right, "#,##0.00" if row.total_sum is not None else "@"),
                (row.average if row.average is not None else "-", align_right, "#,##0.00" if row.average is not None else "@"),
                (row.min_value if row.min_value is not None else "-", align_right, "#,##0.00" if row.min_value is not None else "@"),
                (row.max_value if row.max_value is not None else "-", align_right, "#,##0.00" if row.max_value is not None else "@"),
            ])

        for col_idx, (val, alignment, num_fmt) in enumerate(cols, start=1):
            cell = ws_summary.cell(row=row_cursor, column=col_idx, value=val)
            cell.font = font_tr
            cell.alignment = alignment
            cell.border = cell_border
            if num_fmt != "@":
                cell.number_format = num_fmt
            if c_fill:
                cell.fill = c_fill

        row_cursor += 1

    # Totals Row
    total_cols = [
        ("ИТОГО", align_left, "@"),
        (summary.total_elements, align_right, "#,##0"),
        (summary.present_count, align_right, "#,##0"),
        (summary.present_percent / 100.0, align_right, "0.0%"),
        (summary.missing_count, align_right, "#,##0"),
        (summary.missing_percent / 100.0, align_right, "0.0%"),
    ]
    if is_numeric:
        total_cols.extend([
            (summary.total_sum, align_right, "#,##0.00"),
            (summary.average, align_right, "#,##0.00"),
            (summary.min_value, align_right, "#,##0.00"),
            (summary.max_value, align_right, "#,##0.00"),
        ])

    for col_idx, (val, alignment, num_fmt) in enumerate(total_cols, start=1):
        cell = ws_summary.cell(row=row_cursor, column=col_idx, value=val)
        cell.font = font_total
        cell.fill = fill_total
        cell.alignment = alignment
        cell.border = border_total_top
        if num_fmt != "@":
            cell.number_format = num_fmt

    # Auto-adjust column widths
    for col in ws_summary.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and "\n" not in val_str:
                max_len = len(val_str)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # -------------------------------------------------------------
    # SHEET 2: Элементы БЕЗ атрибута (QA/QC дефекты)
    # -------------------------------------------------------------
    ws_missing = wb.create_sheet(title="Элементы без атрибута")
    ws_missing.views.sheetView[0].showGridLines = True

    ws_missing["A1"] = f"Реестр элементов без атрибута '{summary.attribute.full_path}'"
    ws_missing["A1"].font = font_title
    ws_missing["A2"] = f"Всего пропущенных элементов: {len(summary.missing_elements)}"
    ws_missing["A2"].font = font_subtitle

    missing_headers = ["№", "GlobalId (GUID)", "Класс IFC", "Имя элемента", "Марка / Tag", "Этаж / Уровень", "Тип элемента"]
    m_row = 4
    for col_idx, h_text in enumerate(missing_headers, start=1):
        c = ws_missing.cell(row=m_row, column=col_idx, value=h_text)
        c.font = font_th
        c.fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid") # Dark Red
        c.alignment = align_center
        c.border = cell_border

    m_row += 1
    for idx, item in enumerate(summary.missing_elements, start=1):
        is_even = (idx % 2 == 0)
        c_fill = None if is_even else fill_alt
        row_vals = [
            (idx, align_center),
            (item.get("global_id", ""), align_left),
            (item.get("ifc_class", ""), align_left),
            (item.get("name", ""), align_left),
            (item.get("tag", ""), align_left),
            (item.get("storey", ""), align_left),
            (item.get("element_type", ""), align_left),
        ]
        for col_idx, (val, alignment) in enumerate(row_vals, start=1):
            cell = ws_missing.cell(row=m_row, column=col_idx, value=val)
            cell.font = font_tr
            cell.alignment = alignment
            cell.border = cell_border
            if c_fill:
                cell.fill = c_fill
        m_row += 1

    # Auto-adjust column widths for sheet 2
    for col in ws_missing.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len and "\n" not in val_str:
                max_len = len(val_str)
        ws_missing.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 45)

    # Save workbook
    wb.save(output_path)
    return output_path


def export_to_csv(
    summary: AggregationSummary,
    output_path: str,
    grouping_title: str = "Категория / Группа",
) -> str:
    """
    Export aggregation summary to CSV with UTF-8-SIG for perfect Excel compatibility.
    """
    import csv

    headers = [
        grouping_title,
        "Элементов всего",
        "С атрибутом (шт.)",
        "С атрибутом (%)",
        "Без атрибута (шт.)",
        "Без атрибута (%)",
    ]
    is_numeric = summary.total_sum is not None
    if is_numeric:
        headers.extend(["Сумма", "Среднее", "Минимум", "Максимум"])

    with open(output_path, mode="w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow([f"Отчет по атрибуту: {summary.attribute.full_path}"])
        writer.writerow(headers)

        for row in summary.rows:
            r_data = [
                row.group_name,
                row.total_elements,
                row.present_count,
                f"{row.present_percent}%",
                row.missing_count,
                f"{row.missing_percent}%",
            ]
            if is_numeric:
                r_data.extend([
                    str(row.total_sum).replace(".", ",") if row.total_sum is not None else "",
                    str(row.average).replace(".", ",") if row.average is not None else "",
                    str(row.min_value).replace(".", ",") if row.min_value is not None else "",
                    str(row.max_value).replace(".", ",") if row.max_value is not None else "",
                ])
            writer.writerow(r_data)

        # Total row
        t_data = [
            "ИТОГО",
            summary.total_elements,
            summary.present_count,
            f"{summary.present_percent}%",
            summary.missing_count,
            f"{summary.missing_percent}%",
        ]
        if is_numeric:
            t_data.extend([
                str(summary.total_sum).replace(".", ",") if summary.total_sum is not None else "",
                str(summary.average).replace(".", ",") if summary.average is not None else "",
                str(summary.min_value).replace(".", ",") if summary.min_value is not None else "",
                str(summary.max_value).replace(".", ",") if summary.max_value is not None else "",
            ])
        writer.writerow(t_data)

    return output_path
